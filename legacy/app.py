import os
import json
import io
import re
from datetime import datetime
from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    send_file,
    abort,
)
from flask_sqlalchemy import SQLAlchemy
from flask_login import (
    LoginManager,
    UserMixin,
    login_user,
    logout_user,
    login_required,
    current_user,
)
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv
from google import genai
from google.genai import types

load_dotenv()

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "dev-secret-key-change-me")
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///leadforge.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = "login"
login_manager.login_message = "Please sign in to access that page."
login_manager.login_message_category = "error"

GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY")
GEMINI_MODEL = os.environ.get("GEMINI_MODEL", "gemini-2.5-flash")
LEAD_COUNT = int(os.environ.get("LEAD_COUNT", 50))
gemini_client = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None

VALID_SIZE_BANDS = [
    "1-10",
    "11-50",
    "51-200",
    "201-500",
    "501-1000",
    "1001-5000",
    "5001-10000",
    "10000+",
]

# (header, model attribute, column width) — Status is a blank column
LEAD_COLUMNS = [
    ("S.No", "serial_number", 6),
    ("Company Name", "company_name", 32),
    ("Contact Person", "contact_person", 22),
    ("Designation", "designation", 26),
    ("Contact Email", "contact_email", 34),
    ("Contact LinkedIn URL", "contact_linkedin_url", 40),
    ("Company LinkedIn URL", "company_linkedin_url", 40),
    ("Company Location", "company_location", 26),
    ("Company Size", "company_size", 14),
    ("Status", None, 16),
]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="312E81")  # indigo-900
HEADER_ALIGN = Alignment(vertical="center", horizontal="left")


# ----------------------------------------------------------------------------
# Models
# ----------------------------------------------------------------------------


class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password = db.Column(db.String(255), nullable=False)

    searches = db.relationship(
        "Search",
        backref="user",
        lazy=True,
        cascade="all, delete-orphan",
        order_by="desc(Search.created_at)",
    )


class Search(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    offering = db.Column(db.String(500), nullable=False)
    target_client = db.Column(db.String(500), nullable=False)
    region = db.Column(db.String(200), nullable=False)
    lead_count = db.Column(db.Integer, default=0)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, index=True)

    leads = db.relationship(
        "Lead",
        backref="search",
        lazy=True,
        cascade="all, delete-orphan",
        order_by="Lead.serial_number",
    )


class Lead(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    search_id = db.Column(db.Integer, db.ForeignKey("search.id"), nullable=False)
    serial_number = db.Column(db.Integer)
    company_name = db.Column(db.String(300))
    contact_person = db.Column(db.String(200))
    designation = db.Column(db.String(200))
    contact_email = db.Column(db.String(200))
    contact_linkedin_url = db.Column(db.String(500))
    company_linkedin_url = db.Column(db.String(500))
    company_location = db.Column(db.String(300))
    company_size = db.Column(db.String(50))


@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))


@app.after_request
def add_no_cache_headers(response):
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


# ----------------------------------------------------------------------------
# Auth routes
# ----------------------------------------------------------------------------


@app.route("/")
def index():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")

        user = User.query.filter_by(email=email).first()
        if user and check_password_hash(user.password, password):
            login_user(user)
            return redirect(url_for("dashboard"))

        flash("Invalid email or password.", "error")

    return render_template("login.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")

        if not name or not email or not password:
            flash("All fields are required.", "error")
            return render_template("register.html")
        if len(password) < 6:
            flash("Password must be at least 6 characters.", "error")
            return render_template("register.html")
        if password != confirm:
            flash("Passwords do not match.", "error")
            return render_template("register.html")
        if User.query.filter_by(email=email).first():
            flash("An account with this email already exists.", "error")
            return render_template("register.html")

        new_user = User(
            name=name,
            email=email,
            password=generate_password_hash(password),
        )
        db.session.add(new_user)
        db.session.commit()

        flash("Registration successful. Please sign in.", "success")
        return redirect(url_for("login"))

    return render_template("register.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    flash("You have been signed out.", "success")
    return redirect(url_for("login"))


# ----------------------------------------------------------------------------
# Dashboard + lead generation
# ----------------------------------------------------------------------------


@app.route("/dashboard", methods=["GET"])
@login_required
def dashboard():
    recent = (
        Search.query.filter_by(user_id=current_user.id)
        .order_by(Search.created_at.desc())
        .limit(3)
        .all()
    )
    return render_template(
        "dashboard.html",
        user=current_user,
        recent_searches=recent,
        lead_count=LEAD_COUNT,
    )


def _strip_code_fence(text: str) -> str:
    text = text.strip()
    if text.startswith("```"):
        lines = text.split("\n")
        lines = lines[1:]
        if lines and lines[-1].strip().startswith("```"):
            lines = lines[:-1]
        text = "\n".join(lines).strip()
    return text


def _parse_leads(raw_text: str):
    """Extract a list of lead dicts from a model response that may contain
    citations, markdown fences, or wrapping objects."""
    text = _strip_code_fence(raw_text or "")

    # Prefer the substring between the first '[' and the last ']'
    start = text.find("[")
    end = text.rfind("]")
    if start != -1 and end != -1 and end > start:
        candidate = text[start : end + 1]
        try:
            data = json.loads(candidate)
            if isinstance(data, list):
                return data
        except json.JSONDecodeError:
            pass

    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        return []

    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for key in ("leads", "results", "data", "items"):
            v = data.get(key)
            if isinstance(v, list):
                return v
        for v in data.values():
            if isinstance(v, list):
                return v
    return []


def _normalize_size(value: str) -> str:
    if not value:
        return ""
    v = str(value).strip().replace(",", "").replace(" ", "")
    for band in VALID_SIZE_BANDS:
        if v == band:
            return band
    nums = re.findall(r"\d+", v)
    if not nums:
        return value
    try:
        n = int(nums[-1] if "+" in v else nums[0])
    except ValueError:
        return value
    if n <= 10:
        return "1-10"
    if n <= 50:
        return "11-50"
    if n <= 200:
        return "51-200"
    if n <= 500:
        return "201-500"
    if n <= 1000:
        return "501-1000"
    if n <= 5000:
        return "1001-5000"
    if n <= 10000:
        return "5001-10000"
    return "10000+"


def _build_prompt(offering: str, target_client: str, region: str, count: int) -> str:
    return f"""You are a senior B2B lead-research analyst with live Google Search access.

GOAL
Compile {count} real, currently-operating businesses that fit the buyer profile below. Use Google Search to ground your answers in actual companies that exist on the public web.

BUYER PROFILE
- We offer: {offering}
- We target: {target_client}
- Target region: {region}

REQUIREMENTS
1. PRIORITIZE REAL COMPANIES. Use Google Search to discover companies whose products, services, hiring posts, news, or LinkedIn presence match the target client description in the target region. Do not invent company names.
2. DIVERSIFY across company sizes. Roughly: 30% small (51-500 employees), 40% mid-market (501-5000), 30% large (5001+). Avoid filling the list with only well-known mega-corporations.
3. For each company provide a senior decision-maker whose role would buy this offering (e.g. for HR software → "VP People", "Head of HR", "CHRO"; for cybersecurity → "CISO", "Head of IT Security"; for marketing tools → "VP Marketing", "CMO").
4. Contact email: prefer real emails found in search results. Otherwise construct the most likely business email using the company's verified domain and one of the common patterns: firstname.lastname@domain | firstname@domain | f.lastname@domain. NEVER use @gmail.com, @yahoo.com or other personal providers for contact_email.
5. Contact LinkedIn URL: include a real linkedin.com/in/... URL if you find one in search results; otherwise leave as an empty string "".
6. Company LinkedIn URL: format https://www.linkedin.com/company/<slug> — use the slug you found in search; if uncertain, leave as "".
7. Company location: actual city/region the company operates in (must match or be near the target region).
8. Company size: ONE of exactly these bands — "1-10", "11-50", "51-200", "201-500", "501-1000", "1001-5000", "5001-10000", "10000+".
9. No duplicate company_name entries.

OUTPUT (CRITICAL)
Return ONLY a single JSON array of exactly {count} objects. No markdown, no prose before or after, no code fences, no citations. Each object must have EXACTLY these keys with string values (except serial_number which is an integer):
  "serial_number"            integer 1..{count}
  "company_name"             string
  "contact_person"           string
  "designation"              string
  "contact_email"            string
  "contact_linkedin_url"     string (may be "")
  "company_linkedin_url"     string (may be "")
  "company_location"         string
  "company_size"             string (one of the size bands above)

The first character of your reply must be [ and the last must be ]."""


def _clean_leads(raw_leads: list, count: int) -> list:
    cleaned = []
    seen_companies = set()
    for lead in raw_leads:
        if not isinstance(lead, dict):
            continue
        company = str(lead.get("company_name", "")).strip()
        if not company:
            continue
        key = company.lower()
        if key in seen_companies:
            continue
        seen_companies.add(key)

        cleaned.append(
            {
                "serial_number": int(lead.get("serial_number") or len(cleaned) + 1),
                "company_name": company,
                "contact_person": str(lead.get("contact_person", "")).strip(),
                "designation": str(lead.get("designation", "")).strip(),
                "contact_email": str(lead.get("contact_email", "")).strip(),
                "contact_linkedin_url": str(lead.get("contact_linkedin_url", "")).strip(),
                "company_linkedin_url": str(lead.get("company_linkedin_url", "")).strip(),
                "company_location": str(lead.get("company_location", "")).strip(),
                "company_size": _normalize_size(str(lead.get("company_size", "")).strip()),
            }
        )
        if len(cleaned) >= count:
            break

    for idx, lead in enumerate(cleaned, start=1):
        lead["serial_number"] = idx
    return cleaned


@app.route("/generate", methods=["POST"])
@login_required
def generate():
    offering = request.form.get("offering", "").strip()
    target_client = request.form.get("target_client", "").strip()
    region = request.form.get("region", "").strip()

    if not offering or not target_client or not region:
        flash("Please fill out all three fields.", "error")
        return redirect(url_for("dashboard"))

    if not gemini_client:
        flash(
            "Gemini API key is not configured. Set GEMINI_API_KEY in your .env file.",
            "error",
        )
        return redirect(url_for("dashboard"))

    prompt = _build_prompt(offering, target_client, region, LEAD_COUNT)

    try:
        response = gemini_client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
            config=types.GenerateContentConfig(
                tools=[types.Tool(google_search=types.GoogleSearch())],
                temperature=0.6,
                max_output_tokens=16384,
            ),
        )
        raw_text = response.text or ""
        leads = _clean_leads(_parse_leads(raw_text), LEAD_COUNT)

        if not leads:
            raise ValueError("No leads could be parsed from the model response.")

        search = Search(
            user_id=current_user.id,
            offering=offering,
            target_client=target_client,
            region=region,
            lead_count=len(leads),
        )
        db.session.add(search)
        db.session.flush()

        for lead in leads:
            db.session.add(
                Lead(
                    search_id=search.id,
                    serial_number=lead["serial_number"],
                    company_name=lead["company_name"],
                    contact_person=lead["contact_person"],
                    designation=lead["designation"],
                    contact_email=lead["contact_email"],
                    contact_linkedin_url=lead["contact_linkedin_url"],
                    company_linkedin_url=lead["company_linkedin_url"],
                    company_location=lead["company_location"],
                    company_size=lead["company_size"],
                )
            )
        db.session.commit()
        return redirect(url_for("results", search_id=search.id))

    except Exception as exc:
        db.session.rollback()
        flash(f"Error generating leads: {exc}", "error")
        return redirect(url_for("dashboard"))


# ----------------------------------------------------------------------------
# Results + downloads
# ----------------------------------------------------------------------------


def _owned_search(search_id: int) -> Search:
    search = Search.query.get_or_404(search_id)
    if search.user_id != current_user.id:
        abort(404)
    return search


@app.route("/results/<int:search_id>", methods=["GET"])
@login_required
def results(search_id):
    search = _owned_search(search_id)
    return render_template("results.html", search=search, leads=search.leads)


def _style_header(ws, header_row=1):
    for cell in ws[header_row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[header_row].height = 22
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def _build_single_workbook(search: Search) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    ws.append([h for h, _, _ in LEAD_COLUMNS])
    _style_header(ws)

    for lead in search.leads:
        row = []
        for _, attr, _ in LEAD_COLUMNS:
            if attr is None:
                row.append("")
            else:
                row.append(getattr(lead, attr) or "")
        ws.append(row)

    for i, (_, _, width) in enumerate(LEAD_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    return wb


def _build_all_history_workbook(user: User) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "All Leads"

    meta_headers = ["Search Date", "Offering", "Target Client", "Region"]
    lead_headers = [h for h, _, _ in LEAD_COLUMNS]
    ws.append(meta_headers + lead_headers)
    _style_header(ws)

    searches = (
        Search.query.filter_by(user_id=user.id)
        .order_by(Search.created_at.desc())
        .all()
    )
    for search in searches:
        for lead in search.leads:
            row = [
                search.created_at.strftime("%Y-%m-%d %H:%M"),
                search.offering,
                search.target_client,
                search.region,
            ]
            for _, attr, _ in LEAD_COLUMNS:
                if attr is None:
                    row.append("")
                else:
                    row.append(getattr(lead, attr) or "")
            ws.append(row)

    widths = [18, 32, 32, 22] + [w for _, _, w in LEAD_COLUMNS]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    return wb


def _safe_filename(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_\-]+", "_", text)[:40] or "leads"


@app.route("/download/<int:search_id>", methods=["GET"])
@login_required
def download(search_id):
    search = _owned_search(search_id)
    if not search.leads:
        flash("No leads available to download.", "error")
        return redirect(url_for("results", search_id=search.id))

    wb = _build_single_workbook(search)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    stamp = search.created_at.strftime("%Y%m%d_%H%M")
    fname = f"leadgen_{_safe_filename(search.offering)}_{stamp}.xlsx"
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=fname,
    )


# ----------------------------------------------------------------------------
# History
# ----------------------------------------------------------------------------


@app.route("/history", methods=["GET"])
@login_required
def history():
    searches = (
        Search.query.filter_by(user_id=current_user.id)
        .order_by(Search.created_at.desc())
        .all()
    )
    total_leads = sum(s.lead_count for s in searches)
    return render_template(
        "history.html",
        searches=searches,
        total_leads=total_leads,
    )


@app.route("/history/download", methods=["GET"])
@login_required
def download_all_history():
    if not current_user.searches:
        flash("No search history to download yet.", "error")
        return redirect(url_for("history"))

    wb = _build_all_history_workbook(current_user)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M")
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"leadgen_all_history_{stamp}.xlsx",
    )


@app.route("/history/delete/<int:search_id>", methods=["POST"])
@login_required
def delete_search(search_id):
    search = _owned_search(search_id)
    db.session.delete(search)
    db.session.commit()
    flash("Search deleted.", "success")
    return redirect(url_for("history"))


# ----------------------------------------------------------------------------
# Bootstrap
# ----------------------------------------------------------------------------

with app.app_context():
    db.create_all()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=True)
